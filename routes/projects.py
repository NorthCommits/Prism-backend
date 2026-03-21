import os
import uuid
import traceback
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from pydantic import BaseModel
from typing import Optional, List
from db.supabase import get_supabase
from db.auth import verify_token
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form, Body


router = APIRouter()

# file size limits
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024        # 5MB per file
MAX_PROJECT_SIZE_BYTES = 25 * 1024 * 1024    # 25MB per project
MAX_USER_STORAGE_BYTES = 100 * 1024 * 1024   # 100MB per user

SUPPORTED_FILE_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/csv": "csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/json": "json",
    "text/html": "html",
    "text/x-python": "py",
    "application/x-python-code": "py",
    "text/javascript": "js",
    "text/typescript": "ts",
}

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md", ".csv",
    ".xlsx", ".json", ".html", ".py", ".js", ".ts"
}


class CreateProjectRequest(BaseModel):
    name: str
    description: Optional[str] = None
    instructions: Optional[str] = None
    color: Optional[str] = "#8b5cf6"


class UpdateProjectRequest(BaseModel):
    name: Optional[str] = None
    description: Optional[str] = None
    instructions: Optional[str] = None
    color: Optional[str] = None


class ProjectResponse(BaseModel):
    id: str
    user_id: str
    name: str
    description: Optional[str] = None
    instructions: Optional[str] = None
    color: Optional[str] = "#8b5cf6"
    file_count: Optional[int] = 0
    total_size: Optional[int] = 0
    created_at: str
    updated_at: str


class ProjectFileResponse(BaseModel):
    id: str
    project_id: str
    file_name: str
    file_type: str
    file_size: int
    created_at: str


def get_file_extension(filename: str) -> str:
    """Get file extension from filename."""
    return os.path.splitext(filename)[1].lower()


async def extract_file_content(
    file_bytes: bytes,
    file_type: str,
    filename: str
) -> str:
    """
    Extracts text content from uploaded file.
    Supports: txt, md, csv, json, html, py, js, ts, pdf, docx, xlsx
    """
    ext = get_file_extension(filename)

    try:
        # plain text formats
        if ext in {".txt", ".md", ".py", ".js", ".ts", ".html", ".json", ".csv"}:
            return file_bytes.decode("utf-8", errors="ignore")

        # PDF extraction
        if ext == ".pdf":
            try:
                import io
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                text = ""
                for page in reader.pages:
                    text += page.extract_text() or ""
                return text[:50000]  # limit to 50k chars
            except ImportError:
                return f"[PDF file: {filename} — install pypdf to extract content]"
            except Exception as e:
                return f"[Could not extract PDF content: {e}]"

        # DOCX extraction
        if ext == ".docx":
            try:
                import io
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                text = "\n".join([para.text for para in doc.paragraphs])
                return text[:50000]
            except ImportError:
                return f"[DOCX file: {filename} — install python-docx to extract content]"
            except Exception as e:
                return f"[Could not extract DOCX content: {e}]"

        # XLSX extraction
        if ext == ".xlsx":
            try:
                import io
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
                text = ""
                for sheet in wb.worksheets:
                    text += f"Sheet: {sheet.title}\n"
                    for row in sheet.iter_rows(values_only=True):
                        row_text = "\t".join(
                            str(cell) if cell is not None else ""
                            for cell in row
                        )
                        text += row_text + "\n"
                return text[:50000]
            except Exception as e:
                return f"[Could not extract XLSX content: {e}]"

        return f"[Unsupported file format: {ext}]"

    except Exception as e:
        return f"[Error extracting content: {e}]"


async def get_user_storage_used(user_id: str, client) -> int:
    """Returns total storage used by user in bytes."""
    try:
        response = (
            client.table("project_files")
            .select("file_size")
            .eq("user_id", user_id)
            .execute()
        )
        return sum(f["file_size"] for f in (response.data or []))
    except Exception:
        return 0


async def get_project_storage_used(project_id: str, client) -> int:
    """Returns total storage used by a project in bytes."""
    try:
        response = (
            client.table("project_files")
            .select("file_size")
            .eq("project_id", project_id)
            .execute()
        )
        return sum(f["file_size"] for f in (response.data or []))
    except Exception:
        return 0


# ═══════════════════════════════════════
# PROJECT CRUD
# ═══════════════════════════════════════

@router.get("/projects")
async def list_projects(user_id: str = Depends(verify_token)):
    """List all projects for the user."""
    try:
        client = get_supabase()

        response = (
            client.table("projects")
            .select("*")
            .eq("user_id", user_id)
            .order("updated_at", desc=True)
            .execute()
        )

        projects = response.data or []

        # enrich with file count and total size
        result = []
        for project in projects:
            files_response = (
                client.table("project_files")
                .select("file_size")
                .eq("project_id", project["id"])
                .execute()
            )
            files = files_response.data or []
            project["file_count"] = len(files)
            project["total_size"] = sum(f["file_size"] for f in files)
            result.append(project)

        return result

    except Exception as e:
        print(f"List projects error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/projects")
async def create_project(
    request: CreateProjectRequest,
    user_id: str = Depends(verify_token)
):
    """Create a new project."""
    try:
        client = get_supabase()

        if not request.name or not request.name.strip():
            raise HTTPException(status_code=400, detail="Project name is required")

        response = client.table("projects").insert({
            "user_id": user_id,
            "name": request.name.strip(),
            "description": request.description,
            "instructions": request.instructions,
            "color": request.color or "#8b5cf6"
        }).execute()

        project = response.data[0]
        project["file_count"] = 0
        project["total_size"] = 0

        print(f"Created project: {project['id']} for user {user_id}")
        return project

    except HTTPException:
        raise
    except Exception as e:
        print(f"Create project error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/projects/{project_id}")
async def get_project(
    project_id: str,
    user_id: str = Depends(verify_token)
):
    """Get a specific project with its files."""
    try:
        client = get_supabase()

        response = (
            client.table("projects")
            .select("*")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not response.data:
            raise HTTPException(status_code=404, detail="Project not found")

        project = response.data[0]

        # get files
        files_response = (
            client.table("project_files")
            .select("id, project_id, file_name, file_type, file_size, created_at")
            .eq("project_id", project_id)
            .order("created_at", desc=False)
            .execute()
        )

        project["files"] = files_response.data or []
        project["file_count"] = len(project["files"])
        project["total_size"] = sum(
            f["file_size"] for f in project["files"]
        )

        return project

    except HTTPException:
        raise
    except Exception as e:
        print(f"Get project error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.patch("/projects/{project_id}")
async def update_project(
    project_id: str,
    request: UpdateProjectRequest,
    user_id: str = Depends(verify_token)
):
    """Update project details."""
    try:
        client = get_supabase()

        # verify ownership
        existing = (
            client.table("projects")
            .select("id")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not existing.data:
            raise HTTPException(status_code=404, detail="Project not found")

        update_data = {}
        if request.name is not None:
            update_data["name"] = request.name.strip()
        if request.description is not None:
            update_data["description"] = request.description
        if request.instructions is not None:
            update_data["instructions"] = request.instructions
        if request.color is not None:
            update_data["color"] = request.color

        if not update_data:
            raise HTTPException(status_code=400, detail="No fields to update")

        response = (
            client.table("projects")
            .update(update_data)
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        return response.data[0]

    except HTTPException:
        raise
    except Exception as e:
        print(f"Update project error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/projects/{project_id}")
async def delete_project(
    project_id: str,
    user_id: str = Depends(verify_token)
):
    """Delete a project and all its files."""
    try:
        client = get_supabase()

        # verify ownership
        existing = (
            client.table("projects")
            .select("id")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not existing.data:
            raise HTTPException(status_code=404, detail="Project not found")

        # delete files first
        client.table("project_files").delete().eq(
            "project_id", project_id
        ).execute()

        # unlink conversations
        client.table("conversations").update({
            "project_id": None
        }).eq("project_id", project_id).execute()

        # delete project
        client.table("projects").delete().eq(
            "id", project_id
        ).eq("user_id", user_id).execute()

        print(f"Deleted project {project_id} for user {user_id}")
        return {"success": True}

    except HTTPException:
        raise
    except Exception as e:
        print(f"Delete project error: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ═══════════════════════════════════════
# PROJECT FILES
# ═══════════════════════════════════════

@router.post("/projects/{project_id}/files")
async def upload_project_file(
    project_id: str,
    file: UploadFile = File(...),
    user_id: str = Depends(verify_token)
):
    """Upload a file to a project."""
    try:
        client = get_supabase()

        # verify project ownership
        project = (
            client.table("projects")
            .select("id")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not project.data:
            raise HTTPException(status_code=404, detail="Project not found")

        # validate file extension
        ext = get_file_extension(file.filename or "")
        if ext not in SUPPORTED_EXTENSIONS:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file type: {ext}. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
            )

        # read file
        file_bytes = await file.read()
        file_size = len(file_bytes)

        # check file size
        if file_size > MAX_FILE_SIZE_BYTES:
            raise HTTPException(
                status_code=400,
                detail=f"File too large. Maximum size is 5MB. Your file: {file_size // (1024*1024)}MB"
            )

        if file_size == 0:
            raise HTTPException(status_code=400, detail="File is empty")

        # check project storage limit
        project_used = await get_project_storage_used(project_id, client)
        if project_used + file_size > MAX_PROJECT_SIZE_BYTES:
            raise HTTPException(
                status_code=400,
                detail=f"Project storage limit reached (25MB). Current usage: {project_used // (1024*1024)}MB"
            )

        # check user storage limit
        user_used = await get_user_storage_used(user_id, client)
        if user_used + file_size > MAX_USER_STORAGE_BYTES:
            raise HTTPException(
                status_code=400,
                detail=f"User storage limit reached (100MB). Current usage: {user_used // (1024*1024)}MB"
            )

        # extract text content
        print(f"Extracting content from {file.filename} ({file_size} bytes)")
        content = await extract_file_content(
            file_bytes,
            file.content_type or "",
            file.filename or ""
        )

        # save to database
        result = client.table("project_files").insert({
            "project_id": project_id,
            "user_id": user_id,
            "file_name": file.filename,
            "file_type": ext.lstrip("."),
            "file_size": file_size,
            "file_content": content
        }).execute()

        print(f"Uploaded file {file.filename} to project {project_id}")

        return {
            "id": result.data[0]["id"],
            "project_id": project_id,
            "file_name": file.filename,
            "file_type": ext.lstrip("."),
            "file_size": file_size,
            "created_at": result.data[0]["created_at"]
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"Upload file error: {type(e).__name__}: {e}")
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/projects/{project_id}/files")
async def list_project_files(
    project_id: str,
    user_id: str = Depends(verify_token)
):
    """List all files in a project."""
    try:
        client = get_supabase()

        # verify ownership
        project = (
            client.table("projects")
            .select("id")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not project.data:
            raise HTTPException(status_code=404, detail="Project not found")

        response = (
            client.table("project_files")
            .select("id, project_id, file_name, file_type, file_size, created_at")
            .eq("project_id", project_id)
            .order("created_at", desc=False)
            .execute()
        )

        return response.data or []

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.delete("/projects/{project_id}/files/{file_id}")
async def delete_project_file(
    project_id: str,
    file_id: str,
    user_id: str = Depends(verify_token)
):
    """Delete a file from a project."""
    try:
        client = get_supabase()

        # verify ownership
        file = (
            client.table("project_files")
            .select("id")
            .eq("id", file_id)
            .eq("project_id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not file.data:
            raise HTTPException(status_code=404, detail="File not found")

        client.table("project_files").delete().eq(
            "id", file_id
        ).execute()

        return {"success": True}

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ═══════════════════════════════════════
# CONVERSATION LINKING
# ═══════════════════════════════════════

@router.post("/conversations/{conversation_id}/link-project")
async def link_conversation_to_project(
    conversation_id: str,
    body: dict = Body(...),
    user_id: str = Depends(verify_token)
):
    """Link or unlink a conversation to a project."""
    try:
        project_id = body.get("project_id")
        
        print(f"Linking conversation {conversation_id} to project {project_id}")
        
        client = get_supabase()

        # verify conversation ownership
        conv = (
            client.table("conversations")
            .select("id")
            .eq("id", conversation_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not conv.data:
            raise HTTPException(status_code=404, detail="Conversation not found")

        # if project_id provided, verify ownership
        if project_id:
            project = (
                client.table("projects")
                .select("id")
                .eq("id", project_id)
                .eq("user_id", user_id)
                .execute()
            )
            if not project.data:
                raise HTTPException(status_code=404, detail="Project not found")

        result = client.table("conversations").update({
            "project_id": project_id
        }).eq("id", conversation_id).execute()

        print(f"Update result: {result.data}")

        return {
            "success": True,
            "conversation_id": conversation_id,
            "project_id": project_id
        }

    except HTTPException:
        raise
    except Exception as e:
        print(f"Link project error: {type(e).__name__}: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/projects/{project_id}/conversations")
async def get_project_conversations(
    project_id: str,
    user_id: str = Depends(verify_token)
):
    """Get all conversations linked to a project."""
    try:
        client = get_supabase()

        # verify project ownership
        project = (
            client.table("projects")
            .select("id")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not project.data:
            raise HTTPException(status_code=404, detail="Project not found")

        response = (
            client.table("conversations")
            .select("id, title, created_at, updated_at")
            .eq("project_id", project_id)
            .eq("user_id", user_id)
            .order("updated_at", desc=True)
            .execute()
        )

        return response.data or []

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ═══════════════════════════════════════
# PROJECT CONTEXT FOR CHAT
# ═══════════════════════════════════════

async def get_project_context(project_id: str, user_id: str) -> str:
    """
    Builds context string from project instructions + files.
    Used by chat.py to inject project context into system prompt.
    """
    try:
        client = get_supabase()

        # get project
        project = (
            client.table("projects")
            .select("name, instructions")
            .eq("id", project_id)
            .eq("user_id", user_id)
            .execute()
        )

        if not project.data:
            return ""

        proj = project.data[0]
        context_parts = []

        context_parts.append(f"--- PROJECT CONTEXT: {proj['name']} ---")

        if proj.get("instructions"):
            context_parts.append(
                f"Project Instructions:\n{proj['instructions']}"
            )

        # get file contents
        files = (
            client.table("project_files")
            .select("file_name, file_type, file_content")
            .eq("project_id", project_id)
            .execute()
        )

        if files.data:
            context_parts.append(f"\nProject Files ({len(files.data)} files):")
            for f in files.data:
                if f.get("file_content"):
                    # limit each file to 3000 chars to avoid token overflow
                    content_preview = f["file_content"][:3000]
                    if len(f["file_content"]) > 3000:
                        content_preview += "\n... [content truncated]"
                    context_parts.append(
                        f"\n--- File: {f['file_name']} ---\n{content_preview}"
                    )

        context_parts.append("--- END PROJECT CONTEXT ---")
        return "\n\n".join(context_parts)

    except Exception as e:
        print(f"Get project context error: {e}")
        return ""